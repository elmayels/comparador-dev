from __future__ import annotations

import re
import time
import unicodedata
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


def _txt(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _num(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace("%", "")
    s = s.replace(" ", "")
    if s.count(",") and s.count("."):
        # Colombian/Latam human format: 1.234.567,89 or 1,234,567.89.
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif s.count(",") == 1 and s.count(".") == 0:
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        return float(s)
    except Exception:
        return None


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode("ascii")
    s = re.sub(r"[^a-zA-Z0-9%]+", " ", s).lower().strip()
    return re.sub(r"\s+", " ", s)



def canonical_key(code: str = "", description: str = "") -> str:
    code = str(code or "").strip().lower()
    if code:
        return f"code:{code}"
    desc = str(description or "").strip().lower()
    return "desc:" + " ".join(desc.split())[:120]


def _tokens(s: str) -> set[str]:
    stop = {"de", "del", "la", "el", "los", "las", "y", "en", "con", "para", "por", "a", "un", "una", "m", "ml", "kg", "pza", "und"}
    return {t for t in _norm(s).split() if len(t) > 2 and t not in stop}


@dataclass
class CanonicalConcept:
    code: str = ""
    description: str = ""
    unit: str = ""
    quantity: float | None = None
    unit_price: float | None = None
    amount: float | None = None
    market_unit_price: float | None = None
    market_amount: float | None = None
    market_source: str = ""
    market_state: str = ""
    family: str = ""
    hierarchy_level: int = 0
    original_order: int = 0
    is_executable: bool = False
    weight_pct: float | None = None
    cumulative_pct: float | None = None
    is_pareto_80: bool = False
    source_row: int | None = None


@dataclass
class CanonicalApuItem:
    code: str = ""
    concept_key: str = ""
    description: str = ""
    unit: str = ""
    section: str = ""
    unit_price: float | None = None
    operator: str = "*"
    quantity: float | None = None
    amount: float | None = None
    percent: float | None = None
    market_unit_price: float | None = None
    market_operator: str = ""
    market_quantity: float | None = None
    market_amount: float | None = None
    market_deviation: float | None = None
    matched_reference_code: str = ""
    matched_reference_description: str = ""
    matched_reference_source: str = ""
    # Canonical market provenance flags. These drive the Excel visual cue:
    # market values are bold only when they come from a real reference or differ
    # from the contractor's declared value. Fallback values remain normal.
    market_unit_price_is_fallback: bool = False
    market_operator_is_fallback: bool = False
    market_quantity_is_fallback: bool = False
    market_amount_is_fallback: bool = False
    state: str = ""
    observation: str = ""
    source_row: int | None = None


@dataclass
class CanonicalProvider:
    name: str
    concepts_file: str = ""
    matrix_file: str = ""
    concepts: list[CanonicalConcept] = field(default_factory=list)
    apu_items: list[CanonicalApuItem] = field(default_factory=list)
    # Provider-scoped 80/20. These keys represent the concepts that explain
    # the first 80% of this provider's own catalog amount (P.U. total *
    # quantity). They are intentionally not stored on the shared catalog spine
    # because every provider can have a different Pareto set.
    pareto_concept_keys: list[str] = field(default_factory=list)
    validations: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class CanonicalRun:
    run_id: str
    kind: str
    project_name: str
    providers: list[CanonicalProvider] = field(default_factory=list)
    base_concepts: list[CanonicalConcept] = field(default_factory=list)
    base_apu_items: list[CanonicalApuItem] = field(default_factory=list)
    validations: list[dict[str, Any]] = field(default_factory=list)

    def as_dict(self):
        return asdict(self)


class ReferenceCatalog:
    """Granular Construdata reference loader and matcher.

    The previous implementation relied on ``ws.max_row`` while opening the
    Construdata files in read-only mode. Several of those files report
    ``max_row=None`` (materials, machinery and one labor file), so they were
    silently skipped and only a small labor file was loaded. That is why
    materials almost never matched. This loader reads rows sequentially from
    the first visible sheet and detects the relevant columns from the header.
    """

    MAX_ROWS_PER_FILE = 200000
    MAX_MARKET_OVER_CONTRACTOR_FACTOR = 1.25

    def __init__(self, data_dir: Path):
        self.items: list[dict[str, Any]] = []
        self.index: dict[str, list[dict[str, Any]]] = {}
        self.load(data_dir)

    def _column_map(self, header: list[Any], kind: str) -> dict[str, int | None]:
        normed = [_norm(_txt(h)) for h in header]

        def find(*needles: str) -> int | None:
            for i, h in enumerate(normed):
                if all(n in h for n in needles):
                    return i
            return None

        code_idx = find("codigo") or find("code")
        desc_idx = find("descripcion", "completa") or find("descripcion") or find("concepto")
        unit_idx = find("unidad")
        price_idx = None
        if kind == "MATERIAL":
            price_idx = find("costo")
            # Prefer the final total cost column over intermediate cost columns.
            costo_cols = [i for i, h in enumerate(normed) if h == "costo" or h.endswith(" costo")]
            if costo_cols:
                price_idx = costo_cols[-1]
        elif kind == "MO":
            price_idx = find("costototal") or find("costo", "total") or find("salario", "real")
        elif kind == "MAQUINARIA":
            price_idx = find("costo")
        if price_idx is None:
            price_idx = find("precio") or find("pu")
        return {"code": code_idx, "description": desc_idx, "unit": unit_idx, "price": price_idx}

    def load(self, data_dir: Path):
        for path in data_dir.glob("*.xlsx"):
            lname = path.name.lower()
            if "matrices" in lname:
                continue
            kind = "MATERIAL" if "material" in lname else "MO" if "mano" in lname or "obra" in lname else "MAQUINARIA" if "maquinaria" in lname else "REF"
            try:
                wb = load_workbook(path, read_only=True, data_only=True)
                # Canonical rule: first visible sheet only.
                ws = next((s for s in wb.worksheets if getattr(s, "sheet_state", "visible") == "visible"), wb.worksheets[0])
                rows_iter = ws.iter_rows(values_only=True)
                header = next(rows_iter, None)
                if not header:
                    wb.close()
                    continue
                cmap = self._column_map(list(header), kind)
                loaded = 0
                for row_num, row in enumerate(rows_iter, start=2):
                    if row_num > self.MAX_ROWS_PER_FILE:
                        break
                    vals = list(row)
                    desc_idx = cmap.get("description")
                    price_idx = cmap.get("price")
                    code_idx = cmap.get("code")
                    unit_idx = cmap.get("unit")
                    desc = _txt(vals[desc_idx]) if desc_idx is not None and desc_idx < len(vals) else ""
                    code = _txt(vals[code_idx]) if code_idx is not None and code_idx < len(vals) else ""
                    unit = _txt(vals[unit_idx]) if unit_idx is not None and unit_idx < len(vals) else ""
                    price = _num(vals[price_idx]) if price_idx is not None and price_idx < len(vals) else None
                    if not desc or price is None or price <= 0:
                        continue
                    tokens = _tokens(desc) | _tokens(code)
                    item = {
                        "code": code,
                        "description": desc,
                        "unit": unit,
                        "price": float(price),
                        "kind": kind,
                        "source": path.name,
                        "row": row_num,
                        "tokens": tokens,
                    }
                    self.items.append(item)
                    loaded += 1
                    for tok in item["tokens"]:
                        self.index.setdefault(tok, []).append(item)
                wb.close()
            except Exception:
                continue

    def match(self, description: str, section: str = "", contractor_price: float | None = None, unit: str = "") -> dict[str, Any] | None:
        q = _tokens(description)
        if not q:
            return None
        wanted = ""
        sn = _norm(section)
        if "material" in sn:
            wanted = "MATERIAL"
        elif "mano" in sn or sn == "mo":
            wanted = "MO"
        elif "maquinaria" in sn or "equipo" in sn:
            wanted = "MAQUINARIA"
        best = None
        best_score = 0.0
        candidates: dict[int, dict[str, Any]] = {}
        for tok in q:
            for item in self.index.get(tok, []):
                candidates[id(item)] = item
        iterable = candidates.values() if candidates else self.items
        unit_n = _norm(unit)
        for item in iterable:
            if wanted and item["kind"] != wanted:
                continue
            it = item["tokens"]
            if not it:
                continue
            inter = len(q & it)
            if inter == 0:
                continue
            # Blend coverage over the contractor text with Jaccard to avoid
            # loose one-word matches, and reward unit equality where available.
            coverage = inter / max(len(q), 1)
            jaccard = inter / max(len(q | it), 1)
            score = 0.70 * coverage + 0.30 * jaccard
            if unit_n and _norm(item.get("unit", "")) == unit_n:
                score += 0.08
            if score > best_score:
                best_score = score
                best = item
        if not best or best_score < 0.32:
            return None
        result = {**best, "confidence": round(best_score, 2)}
        if contractor_price is not None and contractor_price > 0 and result["price"] > contractor_price * self.MAX_MARKET_OVER_CONTRACTOR_FACTOR:
            result["rejected"] = True
            result["reject_reason"] = f"Precio Construdata {result['price']:.2f} supera {self.MAX_MARKET_OVER_CONTRACTOR_FACTOR:.2f}x el precio contratista {contractor_price:.2f}"
        return result


def _sheet_score(ws) -> int:
    return min(ws.max_row or 0, 5000) * min(ws.max_column or 0, 50)


def _best_sheet(path: Path, preferred_names: Iterable[str] | None = None):
    """Return the workbook's first visible worksheet.

    Canonical ingestion rule: when an uploaded workbook has multiple tabs, the
    system must process only the first visible tab (Sheet1 or whatever name the
    vendor gave to that first tab). Secondary tabs are ignored unless a future
    explicit workflow asks for cross-sheet validation. This prevents accidental
    mixing of a PU tab with a CATALOGO tab inside the same workbook.

    ``preferred_names`` is kept for backward compatibility with older calls but
    no longer changes the selected sheet.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        if getattr(ws, "sheet_state", "visible") == "visible":
            return wb, ws
    return wb, wb.worksheets[0]


def _hierarchy_level(code: str, executable: bool = False) -> int:
    c = _txt(code)
    if not c:
        return 0
    # 1 -> level 1, 1.1 -> level 2, 1.1.1 -> level 3.
    parts = [p for p in re.split(r"[.]", c) if p != ""]
    return max(1, len(parts)) if parts else (3 if executable else 1)


def _apply_pareto_80(concepts: list[CanonicalConcept]) -> None:
    execs = [c for c in concepts if c.is_executable and (c.amount or 0) > 0]
    total = sum(float(c.amount or 0) for c in execs)
    if total <= 0:
        return
    ordered = sorted(execs, key=lambda c: float(c.amount or 0), reverse=True)
    cumulative = 0.0
    for c in ordered:
        c.weight_pct = float(c.amount or 0) / total
        cumulative += c.weight_pct
        c.cumulative_pct = cumulative
        # Include the concept that crosses the 80% threshold. This is the
        # professional 80/20 selection for negotiation priority.
        if cumulative <= 0.80 or (cumulative - (c.weight_pct or 0)) < 0.80:
            c.is_pareto_80 = True



APU_SECTION_TERMS = {
    "materiales", "mano de obra", "equipo y herramienta", "equipo", "maquinaria",
    "basicos", "básicos", "seccion financiera", "sección financiera",
    "subtotal materiales", "subtotal mano de obra", "subtotal maquinaria",
    "subtotal equipo", "subtotal basicos", "subtotal básicos", "costo directo",
    "costo indirecto", "indirectos", "utilidad", "financiamiento",
    "precio unitario", "total costo unitario", "total por servicio",
}


def _looks_like_apu_matrix_row(code: str, desc: str, unit: str = "", op: str = "") -> bool:
    """Guardrail for Comparativa: catalog parsing must not leak PU/matrix rows.

    Comparativa compares catalog concepts and total P.U. by provider. It must
    never show material/labor/equipment detail rows from the matrix/APU. Those
    rows belong exclusively to Detalle - <Proveedor>.
    """
    ndesc = _norm(desc)
    ncode = _norm(code)
    nunit = _norm(unit)
    nop = _norm(op)
    if not ndesc:
        return False
    if ndesc in APU_SECTION_TERMS or any(term == ndesc for term in APU_SECTION_TERMS):
        return True
    if any(term in ndesc for term in [
        "subtotal materiales", "subtotal mano", "subtotal maquinaria", "subtotal equipo",
        "costo directo", "costo indirecto", "precio unitario", "total costo",
        "seccion financiera", "sección financiera"
    ]):
        return True
    # Common PU detail signals: percentage rows, operators, and section-only rows.
    if ncode.startswith("%") or "%" in code or nop in {"*", "/", "%"} and not re.match(r"^\d+(\.\d+)*$", str(code).strip()):
        if nunit in {"%", "jorn", "hr", "hora", "kg", "m3", "m2", "m", "pza", "l", "lt"} or ncode:
            return True
    if ndesc in {"materiales", "mano de obra", "maquinaria", "equipo", "basicos", "básicos"}:
        return True
    return False


def _detect_header(rows: list[list[Any]], synonyms: dict[str, list[str]]) -> tuple[int, dict[str, int]]:
    best_idx = 0
    best_map: dict[str, int] = {}
    best_score = -1
    for idx, row in enumerate(rows[:25]):
        cells = [_norm(_txt(v)) for v in row]
        colmap: dict[str, int] = {}
        score = 0
        for field, words in synonyms.items():
            for c, text in enumerate(cells):
                if any(w in text for w in words):
                    colmap[field] = c
                    score += 1
                    break
        if score > best_score:
            best_score = score
            best_idx = idx
            best_map = colmap
    return best_idx, best_map


def _header_norms(row: list[Any]) -> list[str]:
    return [_norm(_txt(v)) for v in row]


def _detect_repeated_header_blocks(row: list[Any]) -> dict[str, int]:
    """Detect PMD-style provider + market blocks in a single header row.

    Standard manual PMD matrices usually repeat these headers:
    Código | Concepto | Unidad | P. Unitario | Op. | Cantidad | Importe | % | | P. Unitario | Op. | Cantidad | Importe

    The second repeated price/operator/quantity/amount block is the market
    block even when the word "Mercado" is not present in the same row.
    """
    cells = _header_norms(row)
    colmap: dict[str, int] = {}
    def positions(options: set[str] | tuple[str, ...]) -> list[int]:
        out = []
        for i, text in enumerate(cells):
            if any(opt == text or opt in text for opt in options):
                out.append(i)
        return out
    # provider columns
    for i, text in enumerate(cells):
        if text in {"codigo", "clave"}:
            colmap.setdefault("code", i)
        elif "concepto" in text or "descripcion" in text:
            colmap.setdefault("description", i)
        elif text in {"unidad", "und", "udm"}:
            colmap.setdefault("unit", i)
        elif text in {"op", "operador"}:
            colmap.setdefault("operator", i)
        elif "cantidad" in text or text == "cant":
            colmap.setdefault("quantity", i)
        elif "importe" in text or text in {"total", "valor"}:
            colmap.setdefault("amount", i)
        elif text == "%" or "porcentaje" in text:
            colmap.setdefault("percent", i)
        elif text in {"p unitario", "p u", "precio unitario", "pu"}:
            colmap.setdefault("unit_price", i)
    pu_cols = [i for i, t in enumerate(cells) if t in {"p unitario", "p u", "precio unitario", "pu"}]
    op_cols = [i for i, t in enumerate(cells) if t in {"op", "operador"}]
    qty_cols = [i for i, t in enumerate(cells) if "cantidad" in t or t == "cant"]
    amt_cols = [i for i, t in enumerate(cells) if "importe" in t or t in {"total", "valor"}]
    if len(pu_cols) >= 2:
        colmap["market_unit_price"] = pu_cols[1]
    if len(op_cols) >= 2:
        colmap["market_operator"] = op_cols[1]
    if len(qty_cols) >= 2:
        colmap["market_quantity"] = qty_cols[1]
    if len(amt_cols) >= 2:
        colmap["market_amount"] = amt_cols[1]
    return colmap


def _detect_matrix_header(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    """Detect matrix/APU columns by semantic headers, including PMD variants."""
    # Variant A: standard PMD repeated block in row 2 or nearby.
    best_idx, best_map, best_score = 0, {}, -1
    for idx, row in enumerate(rows[:40]):
        m = _detect_repeated_header_blocks(row)
        score = len([k for k in ["code", "description", "unit", "unit_price", "quantity", "amount"] if k in m]) + 2 * len([k for k in ["market_unit_price", "market_amount"] if k in m])
        if score > best_score:
            best_idx, best_map, best_score = idx, m, score
    if best_score >= 5:
        # Variant B: wide PMD/TAPIAL layout. It often has headers like:
        # Clave / Descripción / Unidad / Cantidad / P. Unitario / Importe / Cantidad / P. Unitario / Importe
        row = rows[best_idx]
        cells = _header_norms(row)
        if len(row) >= 24 and any("clave" == t for t in cells) and any("descripcion" in t for t in cells):
            # Prefer detected positions but override obvious wide-layout market group.
            # 0-based columns observed in TAPIAL: code C, desc F, unit J, qty N, PU S, amount W, market qty Y, market PU Z, market amount AA.
            best_map.update({
                "code": 2, "description": 5, "unit": 9,
                "quantity": 13, "unit_price": 18, "amount": 22,
                "market_quantity": 24, "market_unit_price": 25, "market_amount": 26,
            })
        return best_idx, best_map
    return _detect_header(rows, MATRIX_SYNONYMS)


def _detect_concept_market_columns(header_row: list[Any]) -> dict[str, int]:
    cells = _header_norms(header_row)
    pu_cols = [i for i,t in enumerate(cells) if t in {"p u", "p unitario", "precio unitario", "pu"}]
    amt_cols = [i for i,t in enumerate(cells) if "importe" in t or t in {"subtotal", "total", "valor"}]
    out = {}
    if len(pu_cols) >= 2:
        out["market_unit_price"] = pu_cols[1]
    if len(amt_cols) >= 2:
        out["market_amount"] = amt_cols[1]
    return out


CONCEPT_SYNONYMS = {
    "code": ["codigo", "clave", "partida", "item"],
    "description": ["descripcion", "concepto", "servicio", "actividad"],
    "unit": ["unidad", "und", "udm"],
    "quantity": ["cantidad", "cant", "volumen"],
    "unit_price": ["p u", "precio unitario", "pu", "p unitario"],
    "amount": ["importe", "total", "valor", "monto"],
}

MATRIX_SYNONYMS = {
    "code": ["codigo", "clave", "insumo", "partida"],
    "description": ["descripcion", "concepto", "insumo", "recurso"],
    "unit": ["unidad", "und", "udm"],
    "unit_price": ["p unitario", "precio unitario", "p u", "pu"],
    "operator": ["op", "operador"],
    "quantity": ["cantidad", "rendimiento", "cant"],
    "amount": ["importe", "total", "valor"],
    "percent": ["%", "porcentaje"],
    "market_unit_price": ["mercado p unitario", "mercado precio unitario", "mercado p u", "mercado pu", "referencia p unitario"],
    "market_operator": ["mercado op", "referencia op"],
    "market_quantity": ["mercado cantidad", "mercado cant", "referencia cantidad"],
    "market_amount": ["mercado importe", "mercado total", "referencia importe"],
}


# Recommended percentage quantities/rates derived from construdata_matrices.xlsx.
# These are canonical market-quantity references for percentage rows when no
# equivalent Construdata matrix has been matched yet. The contractor lane remains
# untouched; these values affect only Mercado Cantidad.
CONSTRUDATA_PERCENT_DEFAULTS = {
    "%MO1": {"qty": 0.03, "label": "HERRAMIENTA MENOR", "source": "construdata_matrices default %MO1"},
    "%HERR": {"qty": 0.03, "label": "HERRAMIENTA MENOR", "source": "construdata_matrices default %MO1"},
    "%MO2": {"qty": 0.05, "label": "ANDAMIOS", "source": "construdata_matrices default %MO2"},
    "%MO3": {"qty": 0.05, "label": "MATERIALES MENORES", "source": "construdata_matrices default %MO3"},
    "%MO5": {"qty": 0.02, "label": "EQUIPO DE SEGURIDAD", "source": "construdata_matrices default %MO5"},
    "%EPP": {"qty": 0.02, "label": "EQUIPO DE SEGURIDAD", "source": "construdata_matrices default %MO5"},
}


def _recommended_percent_default(item: "CanonicalApuItem") -> dict[str, Any] | None:
    """Return Construdata recommended percentage rate for known % rows.

    This is intentionally based on code/semantic concept, not on any random
    percent sign inside long descriptions. A service such as BORO-01 may mention
    "10%" in its text, but it is not a percentage row.
    """
    code = str(item.code or "").strip().upper()
    if code in CONSTRUDATA_PERCENT_DEFAULTS:
        return CONSTRUDATA_PERCENT_DEFAULTS[code]
    d = _norm(item.description or "")
    if "herramienta menor" in d:
        return CONSTRUDATA_PERCENT_DEFAULTS["%MO1"]
    if "andamio" in d:
        return CONSTRUDATA_PERCENT_DEFAULTS["%MO2"]
    if "materiales menores" in d or "material menor" in d:
        return CONSTRUDATA_PERCENT_DEFAULTS["%MO3"]
    if "equipo de seguridad" in d or "equipo de proteccion" in d or "proteccion personal" in d or d == "epp":
        return CONSTRUDATA_PERCENT_DEFAULTS["%MO5"]
    return None


def classify_xlsx_role(path: Path) -> str:
    """Classify the workbook's business role by structure, not filename.

    This avoids a common upload problem: a user may upload a PU/APU workbook in
    the concepts slot or a totals-by-concept workbook in the matrix slot. The
    canonical pipeline must decide based on sheet semantics.
    """
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            rows: list[list[Any]] = []
            # Role detection follows the same canonical single-sheet contract:
            # use only the first visible worksheet. Do not scan CATALOGO/PU
            # secondary tabs because upload role resolution is per workbook.
            ws = next((w for w in wb.worksheets if getattr(w, "sheet_state", "visible") == "visible"), wb.worksheets[0])
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 80), values_only=True):
                rows.append(list(row))
            joined = " ".join(_norm(" ".join(_txt(v) for v in r if _txt(v))) for r in rows)
            has_analysis = "analisis" in joined or "analisis de basicos" in joined or "analisis de precios" in joined
            has_sections = any(t in joined for t in ["materiales", "mano de obra", "equipo y herramienta", "precio unitario"])
            has_op = any(_norm(_txt(v)) in {"op", "operador"} for r in rows for v in r)
            has_percent_header = any(_norm(_txt(v)) == "%" or "porcentaje" in _norm(_txt(v)) for r in rows for v in r)
            has_budget = "presupuesto de obra" in joined or "servicios cotizacion" in joined or "servicios cotizacion" in joined
            # PU/APU details have analysis headers, sections and usually Op/% columns.
            if has_analysis and has_sections and (has_op or has_percent_header):
                return "apu_detail"
            # Concept/catalog summaries normally have code/concept/unit/quantity and
            # per-section or total cost columns, but no detailed PU sections.
            if has_budget or ("codigo" in joined and "concepto" in joined and "unidad" in joined and "cantidad" in joined and not has_analysis):
                return "concept_catalog"
            return "unknown"
        finally:
            wb.close()
    except Exception:
        return "unknown"


def parse_concepts(path: Path) -> list[CanonicalConcept]:
    """Parse a contractor/base concept catalog preserving original order and hierarchy.

    Preferred source shape is the CATALOGO sheet used by PU workbooks:
    CODIGO | DESCRIPCION | UNIDAD | CANTIDAD | P U | SUBTOTAL.

    The output intentionally keeps non-executable hierarchy rows so the
    Comparativa can show the catalog exactly as declared, while Pareto 80/20
    is calculated only over executable rows with amount.
    """
    wb, ws = _best_sheet(path, preferred_names=["CATALOGO", "CATÁLOGO", "Catalogo", "CATALOGO DE CONCEPTOS", "Catálogo de conceptos", "Comparativa"])
    try:
        rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 1500), values_only=True)]
        # Guardrail: a PU/APU sheet is not a concept catalog. If we parse it as
        # concepts, insumos and financial rows leak into Comparativa. Let the
        # caller choose another uploaded file as concept source.
        if classify_xlsx_role(path) == "apu_detail":
            return []
        hidx, cmap = _detect_header(rows, CONCEPT_SYNONYMS)
        # If we have the canonical CATALOGO layout but fuzzy detection missed it,
        # scan explicit header names.
        if len(cmap) < 3:
            for idx, row in enumerate(rows[:40]):
                normed = [_norm(_txt(v)) for v in row]
                if any("codigo" == x for x in normed) and any("descripcion" in x for x in normed):
                    hidx = idx
                    cmap = {}
                    for c, text in enumerate(normed):
                        if text == "codigo": cmap["code"] = c
                        elif "descripcion" in text: cmap["description"] = c
                        elif "unidad" in text: cmap["unit"] = c
                        elif "cantidad" in text: cmap["quantity"] = c
                        elif text in {"p u", "pu", "precio unitario"}: cmap["unit_price"] = c
                        elif "subtotal" in text or "importe" in text: cmap["amount"] = c
                    break
        cmarket = _detect_concept_market_columns(rows[hidx] if hidx < len(rows) else [])
        data: list[CanonicalConcept] = []
        current_family = ""
        order = 0
        for ridx, row in enumerate(rows[hidx + 1:], hidx + 2):
            def val(field: str, default_idx: int | None = None):
                idx = cmap.get(field, default_idx)
                return row[idx] if idx is not None and idx < len(row) else None
            code = _txt(val("code", 0))
            desc = _txt(val("description", 1))
            if not desc and not code:
                texts = [_txt(v) for v in row if _txt(v)]
                # Ignore empty exported padding rows.
                if not texts:
                    continue
                desc = max(texts, key=len)
            if not desc or len(desc) < 3:
                continue
            # Ignore catalog grand-total rows. They are financial summary rows,
            # not catalog concepts to be compared or painted in Pareto 80/20.
            if not code and (_norm(desc) in {"subtotal", "i v a", "iva", "total"} or _num(desc) is not None):
                continue
            unit = _txt(val("unit"))
            op_guess = _txt(row[cmap.get("operator", -1)]) if "operator" in cmap and cmap.get("operator", -1) < len(row) else ""
            if _looks_like_apu_matrix_row(code, desc, unit, op_guess):
                continue
            qty = _num(val("quantity"))
            pu = _num(val("unit_price"))
            amount = _num(val("amount"))
            market_pu = _num(row[cmarket.get("market_unit_price")]) if cmarket.get("market_unit_price") is not None and cmarket.get("market_unit_price") < len(row) else None
            market_amount = _num(row[cmarket.get("market_amount")]) if cmarket.get("market_amount") is not None and cmarket.get("market_amount") < len(row) else None
            if amount is None and qty is not None and pu is not None:
                amount = qty * pu
            if pu is None and amount is not None and qty not in (None, 0):
                pu = amount / qty
            if market_amount is None and qty is not None and market_pu is not None:
                market_amount = qty * market_pu
            if market_pu is None and market_amount is not None and qty not in (None, 0):
                market_pu = market_amount / qty
            # Continuation row: PMD catalog descriptions may span multiple visual rows.
            # If a row has no code/unit/amount/PU, append it to the previous concept
            # instead of creating a fake catalog line.
            if (not code) and (not unit) and qty in (None, 0) and pu in (None, 0) and amount in (None, 0) and data and data[-1].is_executable:
                data[-1].description = (data[-1].description + " " + desc).strip()
                continue
            # Executable concept for Comparativa means a catalog/service row
            # explicitly declares unit and quantity > 0. PU/APU detail rows are
            # excluded above and in the writer. Notes, chapters and empty/zero
            # rows do not participate in Comparativa, Pareto or KPIs.
            executable = bool(_norm(unit) not in {"", "nota"} and qty is not None and qty > 0 and (pu is not None or amount is not None))
            level = _hierarchy_level(code, executable)
            if not executable:
                current_family = desc
            order += 1
            data.append(CanonicalConcept(
                code=code,
                description=desc,
                unit=unit,
                quantity=qty,
                unit_price=pu,
                amount=amount,
                market_unit_price=market_pu,
                market_amount=market_amount,
                market_source="Comparativa manual" if market_pu is not None or market_amount is not None else "",
                market_state="Mercado leído de comparativa" if market_pu is not None or market_amount is not None else "",
                family=current_family,
                hierarchy_level=level,
                original_order=order,
                is_executable=executable,
                source_row=ridx,
            ))
        # Final contract guardrail: this is a concept catalog, not PU detail.
        # Keep hierarchy/concept rows, drop any matrix/APU detail that may have
        # leaked in from a wrong sheet or ambiguous workbook.
        data = [c for c in data if not _looks_like_apu_matrix_row(c.code, c.description, c.unit)]
        _apply_pareto_80(data)
        return data[:500]
    finally:
        wb.close()


def _classify_section(text: str, current: str) -> str:
    n = _norm(text)
    if not n:
        return current
    if "material" in n:
        return "MATERIALES"
    if "mano de obra" in n or n == "mo" or "jornal" in n:
        return "MANO DE OBRA"
    if "maquinaria" in n or "equipo" in n or "herramienta" in n:
        return "MAQUINARIA"
    if "financier" in n or "indirect" in n or "utilidad" in n or "costo directo" in n:
        return "SECCION FINANCIERA"
    return current



# -----------------------------------------------------------------------------
# Base budget from engineering concept catalog + Construdata matrices
# -----------------------------------------------------------------------------

@dataclass
class ConstrudataMatrixMatch:
    concept_code: str = ""
    concept_name: str = ""
    concept_description: str = ""
    unit: str = ""
    unit_price: float | None = None
    confidence: float = 0.0
    rows: list[dict[str, Any]] = field(default_factory=list)


def parse_base_concepts(path: Path) -> list[CanonicalConcept]:
    """Parse an engineering/base concept catalog.

    Unlike contractor comparison catalogs, a base concept file may not have a
    contractor P.U. yet. For base-budget generation a valid concept is:
    code + unit + quantity > 0 + description. The P.U./amount are calculated
    later from the matched Construdata matrix.
    """
    wb, ws = _best_sheet(path)
    try:
        rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 3000, 3000), values_only=True)]
        # Detect a compact header like: Part. | Cant. | Uni. | Descripción
        hidx = 0
        cmap: dict[str, int] = {}
        for idx, row in enumerate(rows[:30]):
            cells = [_norm(_txt(v)) for v in row]
            score = 0
            local: dict[str, int] = {}
            for c, t in enumerate(cells):
                if t in {"part", "partida", "codigo", "clave"}:
                    local["code"] = c; score += 1
                elif t in {"cant", "cantidad", "volumen"}:
                    local["quantity"] = c; score += 1
                elif t in {"uni", "unidad", "und", "udm"}:
                    local["unit"] = c; score += 1
                elif "desc" in t or "d e s c" in t or "concepto" in t or "servicio" in t:
                    local["description"] = c; score += 1
                elif t in {"p u", "pu", "precio unitario", "p unitario"}:
                    local["unit_price"] = c
                elif "importe" in t or "subtotal" in t or "total" == t:
                    local["amount"] = c
            if score > len(cmap):
                hidx, cmap = idx, local
        if not {"code", "quantity", "unit", "description"}.issubset(set(cmap)):
            # Fall back to generic parser, then relax executability.
            generic = parse_concepts(path)
            for c in generic:
                c.is_executable = bool(_txt(c.code) and _txt(c.unit) and c.quantity is not None and c.quantity > 0)
            return generic
        concepts: list[CanonicalConcept] = []
        family = ""
        order = 0
        for ridx, row in enumerate(rows[hidx+1:], hidx+2):
            def get(field):
                col = cmap.get(field)
                return row[col] if col is not None and col < len(row) else None
            code = _txt(get("code"))
            desc = _txt(get("description"))
            unit = _txt(get("unit"))
            qty = _num(get("quantity"))
            pu = _num(get("unit_price")) if "unit_price" in cmap else None
            amount = _num(get("amount")) if "amount" in cmap else None
            if not desc and not code:
                continue
            if code and qty is not None and qty > 0 and unit and desc:
                if amount is None and pu is not None:
                    amount = qty * pu
                if pu is None and amount is not None and qty:
                    pu = amount / qty
                executable = True
            else:
                executable = False
                if desc:
                    family = desc
            if not desc:
                continue
            order += 1
            concepts.append(CanonicalConcept(
                code=code,
                description=desc,
                unit=unit,
                quantity=qty,
                unit_price=pu,
                amount=amount,
                family=family,
                hierarchy_level=_hierarchy_level(code, executable),
                original_order=order,
                is_executable=executable,
                source_row=ridx,
            ))
        return concepts
    finally:
        wb.close()


class ConstrudataMatrixCatalog:
    """Index of Construdata matrices for independent base-budget generation.

    The matrix file has no header. The columns used here were inferred from the
    actual file bundled in /data:
    5 code, 6 short name, 7 full concept description, 8 unit, 11 P.U.,
    14 insumo code, 16 full insumo description, 17 insumo unit,
    19 section, 20 insumo P.U., 21 quantity/rendimiento, 24 amount.
    """
    MAX_ROWS = 250000

    def __init__(self, path: Path):
        self.path = path
        self.groups: dict[str, dict[str, Any]] = {}
        self._load()

    def _load(self):
        wb = load_workbook(self.path, read_only=True, data_only=True)
        try:
            ws = next((s for s in wb.worksheets if getattr(s, "sheet_state", "visible") == "visible"), wb.worksheets[0])
            for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                if i > self.MAX_ROWS:
                    break
                vals = list(row)
                if len(vals) < 24:
                    continue
                concept_code = _txt(vals[4])
                concept_name = _txt(vals[5])
                concept_desc = _txt(vals[6])
                concept_unit = _txt(vals[7])
                concept_pu = _num(vals[10])
                item_code = _txt(vals[13])
                item_desc = _txt(vals[15]) or _txt(vals[14])
                item_unit = _txt(vals[16])
                section = _txt(vals[18])
                item_pu = _num(vals[19])
                item_qty = _num(vals[20])
                item_amount = _num(vals[23])
                if not concept_code or not concept_desc or not item_desc:
                    continue
                g = self.groups.setdefault(concept_code, {
                    "code": concept_code,
                    "name": concept_name,
                    "description": concept_desc,
                    "unit": concept_unit,
                    "unit_price": concept_pu,
                    "tokens": _tokens(concept_name + " " + concept_desc),
                    "rows": [],
                })
                g["rows"].append({
                    "code": item_code,
                    "description": item_desc,
                    "unit": item_unit,
                    "section": section,
                    "unit_price": item_pu,
                    "operator": "*",
                    "quantity": item_qty,
                    "amount": item_amount,
                    "source_row": i,
                })
        finally:
            wb.close()

    def match(self, concept: CanonicalConcept) -> ConstrudataMatrixMatch | None:
        q = _tokens((concept.description or "") + " " + (concept.code or ""))
        if not q:
            return None
        unit_n = _norm(concept.unit or "")
        best: dict[str, Any] | None = None
        best_score = 0.0
        for g in self.groups.values():
            it = g["tokens"]
            if not it:
                continue
            inter = len(q & it)
            if inter == 0:
                continue
            coverage = inter / max(len(q), 1)
            jaccard = inter / max(len(q | it), 1)
            score = 0.72 * coverage + 0.28 * jaccard
            if unit_n and _norm(g.get("unit", "")) == unit_n:
                score += 0.14
            # penalize incompatible units, but do not eliminate because many
            # engineering base files use LOTE/PZA while Construdata has a per-unit matrix.
            elif unit_n and g.get("unit"):
                score -= 0.04
            if score > best_score:
                best_score, best = score, g
        if not best or best_score < 0.20:
            return None
        return ConstrudataMatrixMatch(
            concept_code=best["code"],
            concept_name=best["name"],
            concept_description=best["description"],
            unit=best["unit"],
            unit_price=best.get("unit_price"),
            confidence=round(best_score, 2),
            rows=list(best["rows"]),
        )


def _build_base_apu_for_concept(concept: CanonicalConcept, match: ConstrudataMatrixMatch | None) -> list[CanonicalApuItem]:
    key = canonical_key(concept.code, concept.description)
    out: list[CanonicalApuItem] = []
    qty_service = concept.quantity or 0
    if not match:
        out.append(CanonicalApuItem(
            code=concept.code,
            concept_key=key,
            description=f"Análisis base no encontrado: {concept.description}",
            unit=concept.unit,
            section="PARTIDA",
            quantity=qty_service,
            state="Sin matriz Construdata",
            observation="No se encontró matriz equivalente; concepto queda para revisión técnica",
        ))
        return out

    out.append(CanonicalApuItem(
        code=concept.code,
        concept_key=key,
        description=f"Análisis base {concept.code} → {match.concept_code} - {match.concept_name}",
        unit=concept.unit or match.unit,
        section="PARTIDA",
        quantity=qty_service,
        unit_price=match.unit_price,
        amount=(match.unit_price or 0) * qty_service if qty_service else None,
        state="Matriz Construdata",
        observation=f"Match {match.concept_code} · confianza {match.confidence}",
    ))

    section_totals: dict[str, float] = {}
    current_section = ""
    for r in match.rows:
        sec = _canonical_section_name(r.get("section") or "", r.get("description") or "")
        if sec != current_section:
            current_section = sec
            out.append(CanonicalApuItem(concept_key=key, description=sec, section="TÍTULO", state="Sección base", observation="Sección generada desde Construdata"))
        item = CanonicalApuItem(
            code=r.get("code", ""),
            concept_key=key,
            description=r.get("description", ""),
            unit=r.get("unit", ""),
            section=sec,
            unit_price=r.get("unit_price"),
            operator=r.get("operator", "*"),
            quantity=r.get("quantity"),
            amount=r.get("amount"),
            state="Construdata matriz",
            observation=f"{match.concept_code} - {match.concept_name}",
            source_row=r.get("source_row"),
        )
        out.append(item)
        if item.amount is not None:
            section_totals[sec] = section_totals.get(sec, 0.0) + float(item.amount)
    direct = 0.0
    for sec in ["MATERIALES", "MANO DE OBRA", "MAQUINARIA", "BASICOS"]:
        if sec in section_totals:
            subtotal = section_totals[sec]
            out.append(CanonicalApuItem(concept_key=key, description=f"SUBTOTAL {sec}", section=f"SUBTOTAL {sec}", amount=subtotal, state="Subtotal base", observation="Calculado desde filas Construdata"))
            direct += subtotal
    indirect = direct * 0.25
    unit_price = direct + indirect
    concept.unit_price = unit_price
    concept.amount = unit_price * qty_service if qty_service else None
    # Keep the PARTIDA header aligned with the generated base P.U., not the
    # original Construdata matrix P.U. column. The generated base follows our
    # current market financial rule: direct cost + indirect 25%.
    if out:
        out[0].unit_price = unit_price
        out[0].amount = concept.amount
    concept.market_unit_price = unit_price
    concept.market_amount = concept.amount
    concept.market_source = "construdata_matrices.xlsx"
    concept.market_state = f"Match {match.concept_code} · confianza {match.confidence}"
    out.extend([
        CanonicalApuItem(concept_key=key, description="SECCIÓN FINANCIERA", section="TÍTULO", state="Financiero base", observation="Regla canónica de mercado"),
        CanonicalApuItem(concept_key=key, description="COSTO DIRECTO", section="COSTO DIRECTO", amount=direct, state="Financiero base", observation="Suma de subtotales base"),
        CanonicalApuItem(concept_key=key, description="INDIRECTO 25%", section="INDIRECTO", unit_price=direct, operator="*", quantity=0.25, amount=indirect, percent=0.25, state="Financiero base", observation="Indirecto de mercado fijo al 25%"),
        CanonicalApuItem(concept_key=key, description="PRECIO UNITARIO", section="PRECIO UNITARIO", amount=unit_price, state="Financiero base", observation="Costo directo + indirecto 25%"),
        CanonicalApuItem(concept_key=key, description="TOTAL POR SERVICIO", section="TOTAL POR SERVICIO", unit_price=unit_price, operator="*", quantity=qty_service, amount=concept.amount, state="Financiero base", observation="P.U. base × cantidad del catálogo"),
    ])
    return out


def generate_base_budget_from_concepts(concepts: list[CanonicalConcept], data_dir: Path) -> tuple[list[CanonicalApuItem], list[dict[str, Any]]]:
    """Generate a base-budget APU detail from engineering concepts.

    This is the independent base-budget flow: engineering catalog concepts are
    matched against ``construdata_matrices.xlsx`` and rendered through the same
    canonical detail writer used by comparisons. It does not affect comparison
    logic.
    """
    matrix_path = data_dir / "construdata_matrices.xlsx"
    validations: list[dict[str, Any]] = []
    if not matrix_path.exists():
        return [], [{"severity": "Alta", "type": "Base Construdata", "message": "No se encontró data/construdata_matrices.xlsx"}]
    idx = ConstrudataMatrixCatalog(matrix_path)
    items: list[CanonicalApuItem] = []
    executable = [c for c in concepts if c.is_executable and c.quantity is not None and c.quantity > 0 and _txt(c.unit)]
    for c in executable:
        match = idx.match(c)
        if not match:
            validations.append({"severity": "Media", "type": "Match matriz base", "message": f"Sin matriz Construdata para {c.code} - {c.description[:80]}"})
        else:
            validations.append({"severity": "Info", "type": "Match matriz base", "message": f"{c.code} → {match.concept_code} ({match.confidence})"})
        items.extend(_build_base_apu_for_concept(c, match))
    _apply_pareto_80(concepts)
    return items, validations

def _percent_base(description: str, section: str) -> str:
    n = _norm(description + " " + section)
    if "material" in n:
        return "% SOBRE MATERIALES"
    if "mano" in n or "mo" in n or "herramienta" in n or "epp" in n:
        return "% SOBRE MO"
    if "maquinaria" in n or "equipo" in n:
        return "% SOBRE MAQUINARIA"
    if "directo ind" in n or "directo indirect" in n or "financ" in n:
        return "% SOBRE DIRECTO+IND"
    if "directo" in n or "indirect" in n or "utilidad" in n:
        return "% SOBRE DIRECTO"
    return section


def parse_matrix(path: Path, catalog: ReferenceCatalog | None = None) -> list[CanonicalApuItem]:
    wb, ws = _best_sheet(path, preferred_names=["PU", "APU", "MATRIZ", "MATRICES", "Detalle - P1", "Detalle - P2"])
    try:
        rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20000), values_only=True)]
        hidx, cmap = _detect_matrix_header(rows)
        current_section = ""
        current_concept_key = ""
        # PU analysis codes may repeat across chapters. Keep a unique canonical
        # occurrence key so detail blocks do not merge accidentally.
        analysis_key_counts: dict[str, int] = {}
        out: list[CanonicalApuItem] = []

        def cell(row: list[Any], field: str):
            idx = cmap.get(field)
            return row[idx] if idx is not None and idx < len(row) else None

        for ridx, row in enumerate(rows[hidx + 1:], hidx + 2):
            texts = [_txt(v) for v in row if _txt(v)]
            longest = max(texts, key=len) if texts else ""
            # Only update the current APU section from explicit section/header rows.
            # Do not infer a new section from normal item descriptions such as
            # "Equipo de protección" or "Herramienta menor" because that
            # corrupts subsequent subtotals and financial calculations.
            first_norm = _norm(texts[0]) if texts else ""
            section_header_terms = {"materiales", "mano de obra", "equipo y herramienta", "equipo", "maquinaria", "basicos", "básicos", "seccion financiera", "sección financiera"}
            if texts and (first_norm in section_header_terms or (len(texts) <= 2 and any(t in _norm(" ".join(texts)) for t in section_header_terms))):
                current_section = _classify_section(" ".join(texts[:4]), current_section)
            whole = _norm(" ".join(texts))

            if texts and (_norm(texts[0]) in {"partida", "analisis"} or str(texts[0]).strip().lower().startswith(("partida", "análisis", "analisis"))):
                label0 = _norm(texts[0])
                header_code = ""
                header_unit = ""
                header_qty = None
                header_amount = None
                # Typical matrix layout:
                # A: "Análisis:" | B: concept code | D: unit | F: quantity | G: total PU
                if label0.startswith("analisis") and len(row) >= 7:
                    header_code = _txt(row[1])
                    header_unit = _txt(row[3])
                    header_qty = _num(row[5])
                    header_amount = _num(row[6])
                    # Wide/manual PMD layouts can have the label in C and the
                    # analysis code several columns to the right. Use the first
                    # non-empty value after the label when the standard B cell is blank.
                    if not header_code:
                        label_pos = next((i for i, v in enumerate(row) if _norm(_txt(v)).startswith("analisis")), None)
                        if label_pos is not None:
                            for v in row[label_pos + 1:]:
                                txt = _txt(v)
                                if txt and not _norm(txt).startswith(("partida", "analisis")):
                                    header_code = txt
                                    break
                    if not header_unit:
                        # fallback to mapped unit header if available
                        ui = cmap.get("unit")
                        header_unit = _txt(row[ui]) if ui is not None and ui < len(row) else header_unit
                    if header_code:
                        base_key = canonical_key(header_code, "")
                        occurrence = analysis_key_counts.get(base_key, 0) + 1
                        analysis_key_counts[base_key] = occurrence
                        current_concept_key = base_key if occurrence == 1 else f"{base_key}#{occurrence}"
                desc_header = " ".join(texts)
                out.append(CanonicalApuItem(
                    code=header_code,
                    concept_key=current_concept_key,
                    description=desc_header,
                    unit=header_unit,
                    section="PARTIDA",
                    quantity=header_qty,
                    amount=header_amount,
                    state="",
                    observation="Encabezado de análisis detectado",
                    source_row=ridx,
                ))
                continue

            if any(k in whole for k in ["subtotal", "subtotal materiales", "subtotal mano", "subtotal maquinaria", "subtotal equipo", "subtotal basicos", "costo directo", "total costo", "precio unitario", "materiales", "mano de obra", "maquinaria", "equipo y herramienta", "basicos", "seccion financiera", "indirectos", "utilidad"]):
                if len(texts) <= 4 or whole in {"materiales", "mano de obra", "maquinaria", "equipo", "equipo y herramienta", "basicos", "seccion financiera"} or "subtotal" in whole or "total" in whole or "costo directo" in whole or "utilidad" in whole or "precio unitario" in whole:
                    amount_guess = _num(cell(row, "amount"))
                    pct_guess = _num(cell(row, "percent"))
                    nums = [_num(v) for v in row if _num(v) is not None]
                    if amount_guess is None and nums:
                        # Fallback only when the sheet does not expose a mapped Importe column.
                        # Prefer not to take market columns as contractor amount.
                        amount_guess = nums[-2] if len(nums) >= 2 and ("precio" in whole or "total" in whole) else nums[-1]
                    if pct_guess is None and nums:
                        # Financial rows in wide PMD layouts often carry the
                        # percentage in an auxiliary column before the amount.
                        # Prefer the last plausible fraction/percentage-like
                        # number, not necessarily the last numeric cell.
                        candidates = [n for n in nums if abs(n) <= 1]
                        pct_guess = candidates[-1] if candidates else None
                    section = whole.upper()[:28] if ("subtotal" in whole or "total" in whole or "costo" in whole or "utilidad" in whole or "precio" in whole) else "TÍTULO"
                    desc_label = longest.upper()
                    if "subtotal" in whole:
                        target = next((t for t in texts if _norm(t) not in {"subtotal", "subtotal:"} and _num(t) is None), "")
                        target_norm = _norm(target)
                        desc_label = (target if target_norm.startswith("subtotal") else "SUBTOTAL " + target).strip().upper()
                    elif "costo directo" in whole:
                        desc_label = "COSTO DIRECTO"
                    elif "precio unitario" in whole or "total costo" in whole:
                        desc_label = "PRECIO UNITARIO" if "precio" in whole else "TOTAL COSTO UNITARIO"
                    elif "utilidad" in whole:
                        desc_label = "UTILIDAD"
                    mpu = _num(cell(row, "market_unit_price"))
                    mop = _txt(cell(row, "market_operator")) or ""
                    mqty = _num(cell(row, "market_quantity"))
                    mamount = _num(cell(row, "market_amount"))
                    if mamount is None and mpu is not None and mqty is not None:
                        mamount = _calc_amount(mpu, mop, mqty)
                    out.append(CanonicalApuItem(
                        concept_key=current_concept_key,
                        description=desc_label,
                        section=section or "TÍTULO",
                        amount=amount_guess,
                        percent=pct_guess,
                        market_unit_price=mpu,
                        market_operator=mop or "",
                        market_quantity=mqty,
                        market_amount=mamount,
                        state="Mercado declarado en matriz" if mamount is not None or mpu is not None else "",
                        observation="Fila estructural detectada" + (" con mercado declarado" if mamount is not None or mpu is not None else ""),
                        source_row=ridx,
                    ))
                    continue

            desc_idx = cmap.get("description")
            desc = _txt(row[desc_idx]) if desc_idx is not None and desc_idx < len(row) else longest
            if not desc or len(desc) < 3:
                continue
            code = _txt(cell(row, "code"))
            unit = _txt(cell(row, "unit"))
            pu = _num(cell(row, "unit_price"))
            op = _txt(cell(row, "operator")) or "*"
            qty = _num(cell(row, "quantity"))
            amount = _num(cell(row, "amount"))
            pct = _num(cell(row, "percent"))
            mpu = _num(cell(row, "market_unit_price"))
            mop = _txt(cell(row, "market_operator")) or ""
            mqty = _num(cell(row, "market_quantity"))
            mamount = _num(cell(row, "market_amount"))

            if amount is None and pu is not None and qty is not None:
                amount = _calc_amount(pu, op, qty)
            if mamount is None and mpu is not None and mqty is not None:
                # Market calculations must respect the declared contractor operator
                # whenever the market block does not explicitly declare its own operator.
                mamount = _calc_amount(mpu, mop or op, mqty)

            # Concept header rows in a PU matrix define the current concept bucket.
            # All following insumos/subtotals inherit this key until the next header.
            looks_like_concept_header = bool(code and unit and amount is not None and not code.strip().startswith("%") and current_section in {"", "SIN SECCIÓN"})
            if looks_like_concept_header:
                base_key = canonical_key(code, desc)
                occurrence = analysis_key_counts.get(base_key, 0) + 1
                analysis_key_counts[base_key] = occurrence
                current_concept_key = base_key if occurrence == 1 else f"{base_key}#{occurrence}"

            section = current_section or "SIN SECCIÓN"
            is_declared_percent = (op == "%") or unit.strip() == "%" or code.strip().startswith("%")
            if is_declared_percent:
                section = _percent_base(desc, section)
                if qty is None and pct is not None:
                    qty = pct / 100 if pct > 1 else pct

            item = CanonicalApuItem(
                code=code,
                concept_key=current_concept_key,
                description=desc,
                unit=unit,
                section=section,
                unit_price=pu,
                operator=op or "*",
                quantity=qty,
                amount=amount,
                percent=pct,
                market_unit_price=mpu,
                market_operator=mop or "",
                market_quantity=mqty,
                market_amount=mamount,
                market_deviation=(pu / mpu - 1) if pu is not None and mpu else None,
                state="Mercado declarado en matriz" if (mpu is not None or mamount is not None) else "Pendiente mercado",
                observation="Leído desde matriz/APU" + (" · mercado declarado" if (mpu is not None or mamount is not None) else ""),
                source_row=ridx,
            )
            if (item.market_unit_price is None and item.market_amount is None) and catalog:
                match = catalog.match(desc, section, contractor_price=pu, unit=unit)
                if match and not match.get("rejected"):
                    item.market_unit_price = float(match["price"])
                    item.market_quantity = qty
                    item.market_operator = op or "*"
                    item.market_amount = _calc_amount(item.market_unit_price, item.market_operator, item.market_quantity)
                    item.market_deviation = (pu / item.market_unit_price - 1) if pu is not None and item.market_unit_price else None
                    item.matched_reference_code = str(match.get("code", "") or "")
                    item.matched_reference_description = str(match.get("description", "") or "")
                    item.matched_reference_source = str(match.get("source", "") or "")
                    item.market_unit_price_is_fallback = False
                    item.market_operator_is_fallback = True
                    item.market_quantity_is_fallback = True
                    item.market_amount_is_fallback = False
                    item.state = "Match mercado"
                    item.observation = f"{match['source']} fila {match.get('row','')} · confianza {match['confidence']}"
                elif match and match.get("rejected"):
                    # Candidate exists, but it is too expensive compared with the contractor.
                    # Business rule: keep contractor value as market fallback when
                    # Construdata is materially above the contractor price.
                    item.market_unit_price = pu
                    item.market_quantity = qty
                    item.market_operator = op or "*"
                    item.market_amount = amount
                    item.market_deviation = 0 if pu is not None else None
                    item.matched_reference_code = str(match.get("code", "") or "") if match else ""
                    item.matched_reference_description = str(match.get("description", "") or "") if match else ""
                    item.matched_reference_source = str(match.get("source", "") or "") if match else ""
                    item.market_unit_price_is_fallback = True
                    item.market_operator_is_fallback = True
                    item.market_quantity_is_fallback = True
                    item.market_amount_is_fallback = True
                    item.state = "Candidato rechazado - usa contratista"
                    item.observation = f"{match.get('source','Construdata')} fila {match.get('row','')} · {match.get('reject_reason','precio fuera de rango')}"
                else:
                    # Canonical fallback: when Construdata has no usable
                    # reference for an insumo, market columns must still be
                    # populated with the contractor's own cost/price. This keeps
                    # financial subtotals complete and makes the absence of a
                    # reference explicit without blanking the market calculation.
                    item.market_unit_price = pu
                    item.market_quantity = qty
                    item.market_operator = op or "*"
                    item.market_amount = amount
                    item.market_deviation = 0 if pu is not None else None
                    item.market_unit_price_is_fallback = True
                    item.market_operator_is_fallback = True
                    item.market_quantity_is_fallback = True
                    item.market_amount_is_fallback = True
                    item.state = "Sin referencia - usa contratista"
                    item.observation = "Sin match en data; mercado usa valor del contratista"
            out.append(item)
        _post_process_market_financials(out)
        return out[:6000]
    finally:
        wb.close()



def _is_percentage_apu_item(item: CanonicalApuItem) -> bool:
    """True for percentage rows declared inside an APU section.

    These rows are special for the market lane: the contractor values are read
    as-is, but market P.U. must be the accumulated market subtotal of the
    corresponding section before percentage rows, not a catalog unit cost.
    """
    code = str(item.code or "").strip()
    unit = str(item.unit or "").strip()
    op = str(item.operator or "").strip()
    desc = _norm(item.description or "")
    sec = _norm(item.section or "")
    return (
        code.startswith("%")
        or unit == "%"
        or op == "%"
        or sec.startswith("% sobre")
        or any(k in desc for k in ["herramienta menor", "equipo de proteccion", "epp", "%herr"])
    )


def _percentage_base_section(item: CanonicalApuItem) -> str:
    """Return the section subtotal that a percentage row must use as base."""
    sec = _norm(item.section or "")
    desc = _norm(item.description or "")
    joined = f"{sec} {desc}"
    if "material" in joined:
        return "MATERIALES"
    if "mano" in joined or "sobre mo" in joined or "herramienta menor" in joined or "epp" in joined or "proteccion" in joined:
        return "MANO DE OBRA"
    if "maquinaria" in joined or "equipo" in joined:
        return "MAQUINARIA"
    return _canonical_section_name(item.section, item.description)


def _has_reference_market(item: CanonicalApuItem) -> bool:
    """Whether a row has at least one non-fallback market value."""
    return any([
        item.market_unit_price is not None and not item.market_unit_price_is_fallback,
        item.market_quantity is not None and not item.market_quantity_is_fallback,
        item.market_amount is not None and not item.market_amount_is_fallback,
    ])


def _normalize_operator(operator: str | None) -> str:
    op = str(operator or "*").strip()
    if op in {"/", "÷", "div", "DIV"}:
        return "/"
    if op in {"*", "x", "X", "×"}:
        return "*"
    if op == "%":
        return "%"
    return op or "*"


def _calc_amount(unit_price: float | None, operator: str | None, quantity: float | None) -> float | None:
    if unit_price is None or quantity is None:
        return None
    op = _normalize_operator(operator)
    if op == "/":
        return unit_price / quantity if quantity else None
    # Percentage rows are normalized to multiplication before calling this
    # function. Any unknown/blank operator follows the contractor's usual
    # multiplication convention.
    return unit_price * quantity

def _canonical_section_name(section: str, description: str = "") -> str:
    sn = _norm(section or "")
    dn = _norm(description or "")
    n = (sn + " " + dn).strip()
    # Canonical percentage bases declared by the parser have priority over
    # words that may appear in the item description, e.g. EPP contains "Equipo"
    # but is commonly declared as % SOBRE MO.
    if "sobre materiales" in sn:
        return "MATERIALES"
    if "sobre mo" in sn or "mano de obra" in sn:
        return "MANO DE OBRA"
    if "sobre maquinaria" in sn or "maquinaria" in sn:
        return "MAQUINARIA"
    if "material" in n:
        return "MATERIALES"
    # Explicit equipment/maquinaria titles win over the word "herramienta".
    if "maquinaria" in n or "equipo y herramienta" in n or sn == "equipo" or "montacargas" in n:
        return "MAQUINARIA"
    if "mano" in n or sn == "mo" or "herramienta menor" in n or "epp" in n:
        return "MANO DE OBRA"
    if "equipo" in n:
        return "MAQUINARIA"
    if "basico" in n:
        return "BASICOS"
    return section or "SIN SECCIÓN"


def _is_structural_item(item: CanonicalApuItem) -> bool:
    s = item.section or ""
    d = _norm(item.description or "")
    return (
        s in {"PARTIDA", "TÍTULO"}
        or s.startswith("SUBTOTAL")
        or any(k in d for k in ["subtotal", "costo directo", "indirecto", "utilidad", "financiamiento", "precio unitario", "total costo", "total por servicio", "importe", "volumen", "rendimiento"])
    )


def _post_process_market_financials(items: list[CanonicalApuItem]) -> None:
    """Calculate market subtotals and financial rows inside the canonical model.

    Matrix files generally contain contractor PU/APU values only. Market values
    must be calculated from matched granular data and then propagated to:
    section subtotals, volume rows, costo directo, indirectos/utilidad and
    precio unitario. The Excel writer only renders these canonical values.
    """
    groups: dict[str, list[CanonicalApuItem]] = {}
    order: list[str] = []
    for it in items:
        key = it.concept_key or "__global__"
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(it)

    for key in order:
        group = groups[key]
        section_market: dict[str, float] = {}
        section_has_reference: dict[str, bool] = {}
        # Contractor values are not recalculated, but we keep running totals from
        # declared row amounts to identify the base a contractor used for a
        # percentage row. Example: %HERR unit_price 6,342.76 equals the declared
        # MO amount accumulated before the percentage rows.
        contractor_section_running: dict[str, float] = {}
        # Explicit declared totals/subtotals mapped to their market counterpart.
        # Example: consu-mec may use the declared SUBTOTAL MANO DE OBRA as its
        # unit price/base even though it lives in EQUIPO Y HERRAMIENTA.
        declared_base_to_market: list[tuple[float, float, bool, str]] = []
        # For section percentage rows, freeze the base at the first percentage
        # encountered in a target section so sibling percentages share the same
        # base.
        frozen_percent_base_by_section: dict[str, float] = {}
        frozen_percent_ref_by_section: dict[str, bool] = {}
        last_section = ""
        last_importe_market_by_section: dict[str, float] = {}
        subtotal_market_by_section: dict[str, float] = {}
        direct_market: float | None = None
        financial_market_total: float = 0.0

        for it in group:
            dnorm = _norm(it.description or "")
            # Track explicit section titles. Financial labels like
            # "(CI) INDIRECTOS" may have been parsed as TÍTULO by the source
            # layout, but they are calculation rows and must not be skipped.
            if it.section == "TÍTULO" and not any(k in dnorm for k in ["indirect", "utilidad", "financ", "costo directo", "precio unitario", "total costo", "subtotal"]):
                sec = _canonical_section_name(it.description, it.description)
                if sec in {"MATERIALES", "MANO DE OBRA", "MAQUINARIA", "BASICOS"}:
                    last_section = sec
                continue

            # Regular line item: calculate only the market lane. The
            # contractor lane is read and rendered exactly as declared.
            if not _is_structural_item(it):
                sec = _canonical_section_name(it.section, it.description)
                if sec in {"MATERIALES", "MANO DE OBRA", "MAQUINARIA", "BASICOS"}:
                    last_section = sec

                # Percentage rows are not priced like normal insumos. When a
                # contractor declares %HERR, %EPP or any percentage inside a
                # section, its market P.U. must be the accumulated market subtotal
                # of that section before section-percentage rows. This makes
                # multiple percentages in the same block use the same base, e.g.
                # FLEX41.11 %HERR and %EPP both apply over the MO subtotal before
                # those percentages.
                if _is_percentage_apu_item(it):
                    # The target section is where this percentage row contributes
                    # (e.g. %HERR inside MO contributes to MO; consu-mec inside
                    # EQUIPO Y HERRAMIENTA contributes to MAQUINARIA/EQUIPO).
                    target_sec = _canonical_section_name(it.section, it.description)

                    def resolve_base() -> tuple[float, bool, str]:
                        declared_base = it.unit_price
                        # 1) Prefer an explicit prior total/subtotal whose
                        # contractor amount equals the percentage base. This
                        # handles rows such as consu-mec using SUBTOTAL MO as base.
                        if declared_base is not None:
                            for contractor_val, market_val, ref_val, label in reversed(declared_base_to_market):
                                if _values_equal(contractor_val, declared_base, tolerance=0.05):
                                    return market_val, ref_val or not _values_equal(market_val, contractor_val, tolerance=0.05), label
                            # 2) Match the declared running amount of any active
                            # section before percentage rows. This handles %HERR
                            # and %EPP both applying over the MO subtotal before
                            # percentages, not over each other.
                            for sec_name, contractor_val in reversed(list(contractor_section_running.items())):
                                if _values_equal(contractor_val, declared_base, tolerance=0.05):
                                    market_val = section_market.get(sec_name, 0.0)
                                    return market_val, section_has_reference.get(sec_name, False) or not _values_equal(market_val, contractor_val, tolerance=0.05), f"acumulado {sec_name}"
                        # 3) Last safe fallback: use the current target section
                        # market accumulator.
                        market_val = section_market.get(target_sec, 0.0)
                        contractor_val = contractor_section_running.get(target_sec, 0.0)
                        return market_val, section_has_reference.get(target_sec, False) or not _values_equal(market_val, contractor_val, tolerance=0.05), f"acumulado {target_sec}"

                    if target_sec not in frozen_percent_base_by_section:
                        base_val, base_ref, base_label = resolve_base()
                        frozen_percent_base_by_section[target_sec] = base_val
                        frozen_percent_ref_by_section[target_sec] = base_ref
                        it.observation = f"Porcentaje de mercado aplicado sobre {base_label}"
                    else:
                        base_val = frozen_percent_base_by_section.get(target_sec, 0.0)
                        base_ref = frozen_percent_ref_by_section.get(target_sec, False)
                        it.observation = f"Porcentaje de mercado aplicado sobre base congelada de {target_sec}"

                    rec = _recommended_percent_default(it)
                    if rec:
                        qty = float(rec["qty"])
                        qty_ref = True
                        qty_source = str(rec["source"])
                    else:
                        qty = it.market_quantity if it.market_quantity is not None else it.quantity
                        qty_ref = False
                        qty_source = "cantidad contratista"
                    op = "*" if (it.market_operator or it.operator or "*") == "%" else (it.market_operator or it.operator or "*")
                    it.market_unit_price = base_val
                    it.market_operator = op
                    it.market_quantity = qty
                    it.market_amount = _calc_amount(base_val, op, qty)
                    it.market_deviation = (it.unit_price / base_val - 1) if it.unit_price is not None and base_val else None
                    it.market_unit_price_is_fallback = not base_ref
                    it.market_quantity_is_fallback = not qty_ref and _values_equal(qty, it.quantity)
                    it.market_operator_is_fallback = _values_equal(op, it.operator)
                    it.market_amount_is_fallback = (
                        it.market_unit_price_is_fallback
                        and it.market_quantity_is_fallback
                        and _values_equal(it.market_amount, it.amount, tolerance=0.05)
                    )
                    if rec:
                        it.state = "Cantidad mercado recomendada Construdata"
                        it.observation = (it.observation + f" · Mercado Cantidad = {qty:.2%} ({qty_source})").strip()
                    else:
                        it.state = "Cálculo mercado porcentaje" if not it.market_amount_is_fallback else "Sin referencia - usa contratista"
                    if it.market_amount is not None:
                        section_market[target_sec] = section_market.get(target_sec, 0.0) + float(it.market_amount or 0)
                        contractor_section_running[target_sec] = contractor_section_running.get(target_sec, 0.0) + float(it.amount or 0)
                        if base_ref:
                            section_has_reference[target_sec] = True
                    continue

                if it.market_amount is None and it.market_unit_price is not None:
                    qty_for_market = it.market_quantity if it.market_quantity is not None else it.quantity
                    op_for_market = it.market_operator or it.operator or "*"
                    it.market_quantity = qty_for_market
                    it.market_operator = op_for_market
                    it.market_amount = _calc_amount(it.market_unit_price, op_for_market, qty_for_market)
                if it.market_deviation is None and it.unit_price is not None and it.market_unit_price:
                    it.market_deviation = it.unit_price / it.market_unit_price - 1
                if it.market_amount is not None:
                    section_market[sec] = section_market.get(sec, 0.0) + float(it.market_amount or 0)
                    if _has_reference_market(it):
                        section_has_reference[sec] = True
                if it.amount is not None:
                    contractor_section_running[sec] = contractor_section_running.get(sec, 0.0) + float(it.amount or 0)
                continue

            # Structural/calculation rows.
            if "importe" in dnorm and "subtotal" not in dnorm:
                sec = last_section or _canonical_section_name(it.section, it.description)
                val = section_market.get(sec)
                if val is not None:
                    it.market_amount = val
                    it.market_unit_price = val
                    it.market_quantity = None
                    it.market_operator = ""
                    it.state = "Cálculo mercado"
                    it.observation = f"Importe mercado calculado desde insumos de {sec}"
                    last_importe_market_by_section[sec] = val
                    if it.amount is not None and it.market_amount is not None:
                        declared_base_to_market.append((float(it.amount or 0), float(it.market_amount or 0), _has_reference_market(it) or not _values_equal(it.market_amount, it.amount, tolerance=0.05), it.description or "Importe"))
                continue

            if "volumen" in dnorm or "rendimiento" in dnorm:
                sec = last_section or _canonical_section_name(it.section, it.description)
                base = last_importe_market_by_section.get(sec, section_market.get(sec))
                vol = it.quantity
                if base is not None and vol is not None:
                    it.market_unit_price = base
                    # Manual PMD uses both patterns: Volumen multiplies, Rendimiento divides.
                    if "rendimiento" in dnorm:
                        it.market_operator = "/"
                        it.market_amount = base / vol if vol else None
                        it.observation = f"Rendimiento mercado = importe {sec} / rendimiento"
                    else:
                        it.market_operator = "*"
                        it.market_amount = base * vol
                        it.observation = f"Volumen mercado = importe {sec} × volumen"
                    it.market_quantity = vol
                    it.state = "Cálculo mercado"
                    subtotal_market_by_section[sec] = it.market_amount
                    if it.amount is not None and it.market_amount is not None:
                        declared_base_to_market.append((float(it.amount or 0), float(it.market_amount or 0), not _values_equal(it.market_amount, it.amount, tolerance=0.05), it.description or "Volumen/Rendimiento"))
                continue

            if dnorm.startswith("subtotal1") or dnorm.startswith("subtotal2"):
                base = direct_market or 0.0
                val = base + financial_market_total
                if val:
                    it.market_amount = val
                    it.market_unit_price = val
                    it.market_operator = ""
                    it.market_quantity = None
                    it.state = "Cálculo mercado"
                    it.observation = "Subtotal financiero mercado = costo directo + cargos acumulados"
                continue

            if "subtotal" in dnorm:
                sec = _canonical_section_name(it.description, it.description)
                val = subtotal_market_by_section.get(sec, section_market.get(sec))
                if val is not None:
                    it.market_amount = val
                    it.market_unit_price = val
                    it.market_operator = ""
                    it.market_quantity = None
                    it.state = "Cálculo mercado"
                    it.observation = f"Subtotal mercado calculado para {sec}"
                    subtotal_market_by_section[sec] = val
                    if it.amount is not None and it.market_amount is not None:
                        declared_base_to_market.append((float(it.amount or 0), float(it.market_amount or 0), not _values_equal(it.market_amount, it.amount, tolerance=0.05), it.description or "Subtotal"))
                continue

            if "costo directo" in dnorm:
                direct_market = sum(float(v or 0) for v in subtotal_market_by_section.values())
                if direct_market:
                    it.market_amount = direct_market
                    it.market_unit_price = direct_market
                    it.market_operator = ""
                    it.market_quantity = None
                    it.state = "Cálculo mercado"
                    it.observation = "Costo directo mercado = suma de subtotales de secciones"
                    if it.amount is not None and it.market_amount is not None:
                        declared_base_to_market.append((float(it.amount or 0), float(it.market_amount or 0), not _values_equal(it.market_amount, it.amount, tolerance=0.05), it.description or "Costo directo"))
                continue

            if "indirect" in dnorm or "utilidad" in dnorm or "financ" in dnorm:
                base = direct_market
                pct = it.quantity if it.quantity is not None else it.percent
                if base is not None:
                    if "indirect" in dnorm:
                        # Canonical rule: when the contractor declares an
                        # indirect cost row, the market lane must always use
                        # Construdata's standard 25%, regardless of the
                        # contractor's declared percentage.
                        factor = 0.25
                        qty_ref = True
                        obs = "Indirecto mercado estándar Construdata = 25% sobre costo directo"
                    elif pct is not None:
                        factor = pct / 100 if pct > 1 else pct
                        qty_ref = False
                        obs = "Cargo financiero mercado calculado sobre costo directo"
                    else:
                        continue
                    it.market_unit_price = base
                    it.market_operator = "*"
                    it.market_quantity = factor
                    it.market_amount = _calc_amount(base, it.market_operator, factor)
                    it.market_unit_price_is_fallback = False
                    it.market_operator_is_fallback = False
                    it.market_quantity_is_fallback = not qty_ref and _values_equal(factor, it.quantity)
                    it.market_amount_is_fallback = False
                    it.state = "Cálculo mercado"
                    it.observation = obs
                    financial_market_total += float(it.market_amount or 0)
                continue

            if "precio unitario" in dnorm or "total costo" in dnorm or "total por servicio" in dnorm:
                base = direct_market or 0.0
                total = base + financial_market_total
                if total:
                    it.market_amount = total
                    it.market_unit_price = total
                    it.market_operator = ""
                    it.market_quantity = None
                    it.state = "Cálculo mercado"
                    it.observation = "Precio unitario mercado = costo directo + cargos financieros"
                continue



def _values_equal(a: Any, b: Any, *, tolerance: float = 1e-7) -> bool:
    if a in (None, "") and b in (None, ""):
        return True
    an = _num(a)
    bn = _num(b)
    if an is not None and bn is not None:
        return abs(an - bn) <= tolerance
    return str(a or "").strip() == str(b or "").strip()


def _market_is_reference_value(item: CanonicalApuItem, field: str) -> bool:
    """Return True when a market field should be emphasized in Excel.

    Canonical rule: fallback market values are merely contractor values copied to
    avoid blanks and must not be emphasized. Reference values, declared market
    values, or calculated market values that differ from the contractor are
    emphasized. This keeps style driven by the canonical row semantics, not by a
    hardcoded Excel patch.
    """
    if field == "unit_price":
        mv, cv, fb = item.market_unit_price, item.unit_price, item.market_unit_price_is_fallback
    elif field == "operator":
        mv, cv, fb = item.market_operator, item.operator, item.market_operator_is_fallback
    elif field == "quantity":
        mv, cv, fb = item.market_quantity, item.quantity, item.market_quantity_is_fallback
    elif field == "amount":
        mv, cv, fb = item.market_amount, item.amount, item.market_amount_is_fallback
    else:
        return False
    if mv in (None, ""):
        return False
    if fb and _values_equal(mv, cv):
        return False
    return (not fb) or (not _values_equal(mv, cv))

def canonical_rows_from_items(items: list[CanonicalApuItem], include_market: bool = True) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    # If the parser already preserved PU structural rows (PARTIDA, TÍTULO,
    # SUBTOTAL, TOTAL), do not synthesize extra section titles. This keeps the
    # output close to the contractor's PU definition.
    has_structural = any(
        (i.section in {"PARTIDA", "TÍTULO"}) or i.section.startswith("SUBTOTAL") or "TOTAL" in i.section or "COSTO" in i.section
        for i in items
    )
    last_section = None
    for item in items:
        structural = (item.section in {"PARTIDA", "TÍTULO"}) or item.section.startswith("SUBTOTAL") or "TOTAL" in item.section or "COSTO" in item.section or "UTILIDAD" in item.section or "PRECIO" in item.section
        if (not has_structural) and item.section != last_section and item.section and not item.section.startswith("%"):
            rows.append({"code":"", "concept": item.section, "unit":"", "section":"TÍTULO", "pu":"", "op":"", "qty":"", "amount":"", "pct":"", "mpu":"", "mop":"", "mqty":"", "mamount":"", "dev":"", "state":"", "obs":""})
            last_section = item.section
        amount = item.amount if item.amount is not None else ""
        market_amount = item.market_amount if item.market_amount is not None else ""
        rows.append({
            "code": item.code,
            "concept": item.description,
            "unit": item.unit,
            "section": item.section,
            "pu": "" if structural else item.unit_price if item.unit_price is not None else "",
            "op": "" if structural else item.operator,
            "qty": "" if structural else item.quantity if item.quantity is not None else "",
            "amount": amount,
            "pct": item.percent / 100 if item.percent and item.percent > 1 else item.percent if item.percent is not None else "",
            "mpu": item.market_unit_price if include_market and item.market_unit_price is not None else "",
            "mop": item.market_operator if include_market and (item.market_unit_price is not None or item.market_amount is not None) else "",
            "mqty": item.market_quantity if include_market and item.market_quantity is not None else "",
            "mamount": market_amount if include_market else "",
            "match_ref": (f"{item.matched_reference_code} - {item.matched_reference_description}" if item.matched_reference_code and item.matched_reference_description else item.matched_reference_description or item.matched_reference_code or "") if include_market else "",
            "dev": item.market_deviation if include_market and item.market_deviation is not None else "",
            "mpu_ref": _market_is_reference_value(item, "unit_price") if include_market else False,
            "mop_ref": _market_is_reference_value(item, "operator") if include_market else False,
            "mqty_ref": _market_is_reference_value(item, "quantity") if include_market else False,
            "mamount_ref": _market_is_reference_value(item, "amount") if include_market else False,
            "state": item.state,
            "obs": item.observation,
        })
    return rows



def _concept_apu_link_map(provider: CanonicalProvider) -> dict[str, str]:
    """Map APU analysis keys to catalog concept keys.

    Primary link is exact key/code. When vendor files use different coding
    systems (e.g. catalog 1.1.1 but PU analysis BS.01), create a high-confidence
    order-based bridge only between executable catalog concepts and detected APU
    analysis headers. The bridge is part of the canonical model flow so
    Comparativa can receive market P.U. from the correct PU block without
    depending on matching codes.
    """
    exec_concepts = [c for c in provider.concepts if c.is_executable]
    concept_keys = [canonical_key(c.code, c.description) for c in exec_concepts]
    direct_keys = set(concept_keys)
    apu_header_items = []
    seen = set()
    for it in provider.apu_items:
        key = it.concept_key or canonical_key(it.code, it.description)
        if not key or key in seen:
            continue
        # Only real analysis headers participate in order-linking. Ignore
        # broader chapter/partida title rows with no analysis code because they
        # throw off the one-to-one CAT ↔ PU order bridge.
        if it.section == "PARTIDA" and it.code:
            seen.add(key)
            apu_header_items.append(it)
        elif it.code and key not in direct_keys and it.unit and (it.amount is not None or it.quantity is not None):
            seen.add(key)
            apu_header_items.append(it)
    apu_keys = [it.concept_key or canonical_key(it.code, it.description) for it in apu_header_items]
    mapping: dict[str, str] = {k: k for k in apu_keys if k in direct_keys}

    # Order bridge: use it when both sides have the same cardinality. This is
    # the common PMD/manual case where CAT uses 1.1.1, 1.1.2... and PU uses
    # BS.01, BS.02... in the same declared order.
    if len(exec_concepts) == len(apu_keys) and len(exec_concepts) > 0:
        confident = 0
        for c, it in zip(exec_concepts, apu_header_items):
            checks = 0
            passed = 0
            if c.unit and it.unit:
                checks += 1
                passed += 1 if _norm(c.unit) == _norm(it.unit) else 0
            if c.quantity is not None and it.quantity is not None:
                checks += 1
                passed += 1 if _values_equal(c.quantity, it.quantity, tolerance=0.01) else 0
            # Header amount may be the P.U. rather than catalog importe, so accept
            # either unit price or amount similarity as supporting evidence.
            if it.amount is not None:
                checks += 1
                passed += 1 if (_values_equal(c.unit_price, it.amount, tolerance=0.05) or _values_equal(c.amount, it.amount, tolerance=0.05)) else 0
            if checks == 0 or passed >= max(1, min(2, checks)):
                confident += 1
        # Do not require every row to have comparable fields; messy vendor files
        # often omit unit/quantity in the header. If most rows pass, link by order.
        if confident >= max(1, int(len(exec_concepts) * 0.65)):
            for c, apu_key in zip(exec_concepts, apu_keys):
                mapping[apu_key] = canonical_key(c.code, c.description)
            provider.validations.append({
                "severity": "Info",
                "type": "Vínculo catálogo-matriz",
                "message": f"Se vinculó catálogo con PU por orden/unidad/cantidad para {len(exec_concepts)} conceptos porque los códigos no coinciden."
            })
    return mapping


def apply_provider_market_to_concepts(provider: CanonicalProvider) -> None:
    """Attach market totals to provider concepts from its canonical APU detail.

    This is a canonical model step, not an Excel styling patch. Comparativa
    compares concept-level P.U. totals; Detalle explains the APU. Therefore the
    provider must expose concept-level market values derived from the APU market
    rows or from granular reference matches. When catalog and PU use different
    codes, a canonical ConceptApuLink-like bridge maps PU analysis keys to
    catalog concept keys by order/unit/quantity/amount.
    """
    if not provider.concepts or not provider.apu_items:
        return
    concept_map = {canonical_key(c.code, c.description): c for c in provider.concepts}
    apu_to_concept = _concept_apu_link_map(provider)
    direct: dict[str, CanonicalApuItem] = {}
    aggregates: dict[str, float] = {}
    for item in provider.apu_items:
        raw_key = item.concept_key or canonical_key(item.code, item.description)
        key = apu_to_concept.get(raw_key, raw_key)
        if not key or key not in concept_map:
            continue
        dnorm = _norm(item.description or "")
        # Preferred source: canonical financial row "PRECIO UNITARIO". This is
        # the market P.U. for the catalog concept. The Comparativa then multiplies
        # it by catalog quantity for market importe.
        if ("precio unitario" in dnorm or "total costo unitario" in dnorm) and item.market_amount is not None:
            direct[key] = item
            continue
        # Secondary source: concept header rows if a matrix explicitly declares
        # market values there.
        if item.code and (apu_to_concept.get(canonical_key(item.code, item.description), canonical_key(item.code, item.description)) == key) and (item.market_amount is not None or item.market_unit_price is not None):
            direct[key] = item
        # Aggregate only true line-level market amounts as fallback. Exclude all
        # structural/calculation rows to avoid double counting subtotals, volume,
        # costo directo and precio unitario.
        structural = _is_structural_item(item)
        if not structural and item.market_amount is not None:
            aggregates[key] = aggregates.get(key, 0.0) + float(item.market_amount or 0)
    for key, concept in concept_map.items():
        qty = concept.quantity or 1
        if key in direct:
            item = direct[key]
            # For a final precio unitario row, market_amount represents market P.U.
            mpu = item.market_amount if item.market_amount is not None else item.market_unit_price
            mamount = (float(mpu) * float(qty)) if mpu is not None and qty else mpu
            concept.market_unit_price = mpu
            concept.market_amount = mamount
            concept.market_source = "APU financiero"
            concept.market_state = "Mercado calculado desde detalle"
        elif key in aggregates:
            mpu = aggregates[key]
            concept.market_unit_price = mpu
            concept.market_amount = (mpu * float(qty)) if qty else mpu
            concept.market_source = "APU granular"
            concept.market_state = "Mercado agregado desde insumos"

def run_id(prefix: str) -> str:
    return f"{prefix}-{int(time.time())}"
